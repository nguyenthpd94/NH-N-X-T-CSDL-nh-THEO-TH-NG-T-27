
import streamlit as st
import google.generativeai as genai
from PIL import Image
import tempfile
import io
import pandas as pd
from docx import Document
from openpyxl.utils import get_column_letter

# ================== CẤU HÌNH TRANG ==================
st.set_page_config(
    page_title="Trợ lý nhận xét TT27 theo điểm số",
    page_icon="💎",
    layout="centered"
)

# ================== CSS ==================
st.markdown("""
<style>
[data-testid="stAppViewContainer"] { background-color: #f4f6f9; }
.header-box {
    background: linear-gradient(135deg,#667eea,#764ba2);
    padding:25px;border-radius:15px;color:white;text-align:center;
}
div.stButton > button {
    background: linear-gradient(90deg,#667eea,#764ba2);
    color:white;border:none;padding:14px;font-weight:bold;
    border-radius:10px;width:100%;font-size:17px;
}
</style>
""", unsafe_allow_html=True)

# ================== HÀM XỬ LÝ ==================
def score_level(score):
    try:
        s = float(score)
    except:
        return None
    if s >= 9: return "9-10"
    if s >= 8: return "8"
    if s >= 7: return "7"
    if s >= 6: return "6"
    if s >= 5: return "5"
    return "<5"

def clean_comment(text):
    if not text:
        return ""
    text = text.strip().lstrip("-•* ")
    return text[0].upper() + text[1:] if len(text) > 1 else text

def extract_comments_by_score(text):
    pools = {}
    current = None
    for line in text.split("\\n"):
        line = line.strip()
        if not line:
            continue
        if line.startswith("###"):
            current = line.replace("###", "").replace("MỨC ĐIỂM", "").strip()
            pools[current] = []
            continue
        if line.startswith("-") and current:
            pools[current].append(clean_comment(line[1:].strip()))
    return pools

# ================== GIAO DIỆN ==================
st.markdown("""
<div class="header-box">
<h2>💎 TRỢ LÝ NHẬN XÉT TIỂU HỌC THEO ĐIỂM SỐ (TT27)</h2>
<p>Tác giả: Nguyễn Văn Nguyên</p>
</div>
""", unsafe_allow_html=True)

# ================== API KEY ==================
with st.sidebar:
    st.header("🔐 Cấu hình API")
    default_key = st.secrets.get("GEMINI_API_KEY", "")
    manual_key = st.text_input("Nhập API Key:", type="password")
    api_key = manual_key or default_key

if api_key:
    genai.configure(api_key=api_key)
else:
    st.warning("⚠️ Chưa có API Key")

# ================== INPUT ==================
student_file = st.file_uploader("📂 File danh sách học sinh (.xlsx)", type=["xlsx"])
evidence_files = st.file_uploader(
    "📂 Minh chứng (ảnh / PDF / Word – không bắt buộc)",
    type=["png", "jpg", "pdf", "docx"],
    accept_multiple_files=True
)

if student_file:
    df = pd.read_excel(student_file, engine="openpyxl")
    st.dataframe(df.head())

    col_score = st.selectbox("📌 Cột điểm", df.columns)
    col_new = st.text_input("📌 Tên cột nhận xét", "Nhận xét GV")
    mon_hoc = st.text_input("📚 Môn học", "Khoa học")
    chu_de = st.text_input("📝 Bài học", "Chủ đề")

    if st.button("🚀 TẠO NHẬN XÉT THEO ĐIỂM"):
        df["__ScoreLevel__"] = df[col_score].apply(score_level)
        score_counts = df["__ScoreLevel__"].value_counts()

        st.write("📊 Phân bố điểm:")
        st.write(score_counts)

        context_text = ""
        media_files = []

        for f in evidence_files:
            if f.name.endswith(".docx"):
                doc = Document(f)
                context_text += "\\n".join(p.text for p in doc.paragraphs)
            elif f.type == "application/pdf":
                with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                    tmp.write(f.getvalue())
                    media_files.append(genai.upload_file(tmp.name))
            else:
                media_files.append(Image.open(f))

        prompt = f"""
Bạn là giáo viên tiểu học. Viết nhận xét học tập môn {mon_hoc}, bài {chu_de}.

QUY TẮC:
- Mỗi nhận xét dùng cho 1 học sinh.
- Không dùng từ: Em, Con, Bạn.
- Không viết in hoa toàn bộ.
- Độ dài 2–3 câu, đúng tinh thần Thông tư 27.
- Nhận xét PHÙ HỢP VỚI ĐIỂM SỐ.

YÊU CẦU SỐ LƯỢNG:
{chr(10).join([f"- {v} nhận xét cho mức điểm {k}" for k,v in score_counts.items()])}

ĐỊNH DẠNG TRẢ VỀ:
### MỨC ĐIỂM 9-10
- ...
### MỨC ĐIỂM 8
- ...
### MỨC ĐIỂM 7
- ...
### MỨC ĐIỂM 6
- ...
### MỨC ĐIỂM 5
- ...
### MỨC ĐIỂM <5
- ...
"""

        try:
            try:
    model = genai.GenerativeModel("models/gemini-pro")
    response = model.generate_content([prompt] + media_files)

            pools = extract_comments_by_score(response.text)

            def assign_comment(row):
                lvl = row["__ScoreLevel__"]
                if lvl in pools and pools[lvl]:
                    return pools[lvl].pop(0)
                return "Hoàn thành nhiệm vụ học tập theo yêu cầu."

            df[col_new] = df.apply(assign_comment, axis=1)
            df.drop(columns="__ScoreLevel__", inplace=True)

            output = io.BytesIO()
            with pd.ExcelWriter(output, engine="openpyxl") as writer:
                df.to_excel(writer, index=False, sheet_name="Data")
                ws = writer.sheets["Data"]
                col_idx = df.columns.get_loc(col_new) + 1
                ws.column_dimensions[get_column_letter(col_idx)].width = 60
            output.seek(0)

            st.success("✅ Hoàn thành!")
            st.download_button(
                "⬇️ Tải file Excel kết quả",
                output,
                "NhanXet_TheoDiem_TT27.xlsx",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )

        except Exception as e:
            st.error(f"Lỗi: {e}")

st.markdown("<div style='text-align:center;color:#888;margin-top:40px;'>© 2026 - Thầy Nguyên</div>", unsafe_allow_html=True)
